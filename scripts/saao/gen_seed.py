#!/usr/bin/env python3
"""One-time generator for the FY2026 General Fund SAAO seed (spec §12, "One-time seed").

Dev-only. Not part of the app and not an import feature: it was run once against the
Budget Office's October 2026 workbook to produce a reviewed migration.

    python3 -I scripts/saao/gen_seed.py "<workbook.xlsx>" \
        > supabase/migrations/20261020000600_saao_seed_fy2026_gf.sql

    python3 -I scripts/saao/gen_seed.py "<workbook.xlsx>" --outline   # print the inferred tree

The workbook is never copied into the repo. The SQL goes to stdout; the self-check report
goes to stderr. The script exits non-zero (and prints nothing to stdout) when any figure,
section total or grand total differs from the workbook's cached values.

Requires openpyxl (tested with 3.1.5).
"""
import re
import sys
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils import get_column_letter as CL

FIRST_ROW, LAST_ROW = 13, 1362        # SAAO body
GRAND_ROW = 1364                       # SAAO grand total
HEADER_ROWS = (3, 4, 5)                # Republic / Province / City
CERTIFIER_ROWS = (1368, 1369)          # name / title (column X)
RR_HEADER_ROW, RR_NOTE_ROW = 10, 9
RR_COLS = [CL(c) for c in range(CI('G'), CI('AF') + 1)]
SB_COLS = {1: 'K', 2: 'L', 3: 'M'}     # SB-4/SB-5 (N, O) are empty this year (checked below)
OBL_MONTH_COLS = [CL(c) for c in range(CI('G'), CI('R') + 1)]   # January..December
EXTRA_COLS = {'cna_1': 'AA', 'cna_2': 'AC', 'cna_3': 'AD', 'savings': 'AE'}
EFFECTIVE_DATE = '2026-01-01'
REMARK = 'Effective date not recorded in the workbook — set the actual date'
SOURCE_NAME = 'GENERAL FUND- OCTOBER  2026.xlsx'

CENT = Decimal('0.01')
CODE_RE = re.compile(r'\d-\d{2}-\d{2}-\d{3}(?:-\d{2,3})?(?:/\d{3})?')
TOTAL_TEXT_RE = re.compile(r'^(sub\s*total|grand\s*total|total)\b', re.I)
CLASS_TOTAL_RE = re.compile(
    r'^(grand\s+)?total\s+(other\s+)?(mooe|maintenance|personal\s+services|capital\s+outlay)', re.I)
CLASS_NAMES = {'PS': 'PERSONAL SERVICES', 'MOOE': 'MAINTENANCE AND OTHER OPERATING EXPENSES',
               'CO': 'CAPITAL OUTLAY'}

# ---- Hand decisions for this one workbook (documented in the task-7 report) ----------------
# Section 425 has no heading of its own: its rows are "MAINTENANCE AND OTHER OPERATING
# EXPENSES", then "SENIOR CITIZEN" and "PERSON WITH DISABILITY" side by side. Its label is
# taken from its total row "TOTAL APPROPRIATIONS -Senior & Person w/ Disability".
SECTION_LABEL_FROM_TOTAL = {425}
# Rows whose kind the generic rules get wrong (figure-free rows only; figures never move
# between rows). 'group' = heading with children.
KIND_OVERRIDES = {
    1325: 'group',   # "5% Local DRRM 2025- SB3": heads rows 1326-1345 (no figures of its own)
    1326: 'group',   # "Development of Science-based Mainstreaming Tools ...": heads MOOE/CO 1327-1345
}
# Rows the indentation places under the wrong heading: row -> parent row (same block).
PARENT_OVERRIDES = {
    1318: 1252,      # "5% Local DRRM 2024- SB3" is a continuing fund beside the 70% Management Fund,
    1325: 1252,      # "5% Local DRRM 2025- SB3" likewise,
    1347: 1252,      # and "QUICK RESPONSE FUND ( 30% )" is the other 30% of the calamity fund.
}
# Rows typed at the wrong indent among their siblings: row -> indent column (A=0).
INDENT_OVERRIDES = {
    120: 2,          # "Local Governance Performance Mgmt. System" sits with PMC, Zoning, ... (column C)
    126: 3,          # "Office Supplies" under Local Shelter Planning, like rows 125 and 127 (column D)
}


def die(msg):
    print('ERROR: ' + msg, file=sys.stderr)
    sys.exit(1)


def collapse(s):
    return re.sub(r'\s+', ' ', str(s)).strip()


def dec(v):
    """Cached cell value -> Decimal pesos (2 dp). None/blank -> 0. Errors raise."""
    if v is None or v == '':
        return Decimal(0)
    if isinstance(v, bool):
        raise ValueError(f'boolean {v!r}')
    if isinstance(v, (int, float)):
        return Decimal(repr(v)).quantize(CENT, ROUND_HALF_UP)
    if isinstance(v, str) and v.strip() == '':
        return Decimal(0)
    raise ValueError(f'not a number: {v!r}')


def numz(v):
    """Like dec(), but text that is not a number (e.g. a stray '.') counts as 0."""
    try:
        return dec(v)
    except ValueError:
        return Decimal(0)


def is_formula(v):
    return isinstance(v, str) and v.startswith('=')


def half(x):
    """round(x/2, 2), half away from zero (R-B)."""
    q = (abs(x) / 2).quantize(CENT, ROUND_HALF_UP)
    return q if x >= 0 else -q


class Sheet:
    def __init__(self, wf, wv, name):
        self.f = wf[name]
        self.v = wv[name]

    def formula(self, r, col):
        return self.f.cell(r, CI(col)).value

    def value(self, r, col):
        return self.v.cell(r, CI(col)).value


class Node:
    def __init__(self, key, kind, label, parent=None, indent=0, code=None, cls=None,
                 total_label=None, synthetic=False):
        self.key = key                # workbook row number (unique; total row for synthesized nodes)
        self.kind = kind
        self.label = label
        self.parent = parent
        self.children = []
        self.indent = indent
        self.code = code
        self.cls = cls
        self.total_label = total_label
        self.synthetic = synthetic
        self.caps = any(ch.isalpha() for ch in label) and label == label.upper()
        self.child_indent = None      # indent of the first row placed under this heading
        if parent is not None:
            parent.children.append(self)

    def depth(self):
        d, p = 0, self.parent
        while p is not None:
            d, p = d + 1, p.parent
        return d

    def lines(self):
        if self.kind == 'line':
            yield self
        for c in self.children:
            yield from c.lines()


def classification_of(text):
    u = collapse(text).upper().rstrip(':').strip()
    if u in ('PERSONAL SERVICES',):
        return 'PS'
    if u in ('MAINTENANCE AND OTHER OPERATING EXPENSES', 'MOOE') or \
            u.startswith('MAINTENANCE AND OTHER OPERATING') or u.startswith('MAINTENANCE & OTHER OPERATING'):
        return 'MOOE'
    if u == 'CAPITAL OUTLAY':
        return 'CO'
    return None


def class_of_total(text):
    m = CLASS_TOTAL_RE.match(text or '')
    if not m:
        return None
    w = m.group(3).lower()
    return 'PS' if w.startswith('personal') else 'CO' if w.startswith('capital') else 'MOOE'


class Seed:
    def __init__(self, path):
        wf = openpyxl.load_workbook(path)
        wv = openpyxl.load_workbook(path, data_only=True)
        self.S = Sheet(wf, wv, 'SAAO')
        self.O = Sheet(wf, wv, 'OBLIGATION')
        self.R = Sheet(wf, wv, 'RealignmentReversion')
        self.errors = []
        self.warnings = []
        self.nodes = {}               # key -> Node
        self.order = []               # creation order
        self.owners = {}              # total row -> node that carries its total
        self.skipped_extras = set()   # (row, field) cells that are formula errors, not tags

    # ---------------------------------------------------------------- row facts
    def err(self, msg):
        self.errors.append(msg)

    def text(self, r):
        for c in range(1, 7):
            v = self.S.f.cell(r, c).value
            if v is not None and collapse(v):
                return collapse(v), c - 1
        return None, None

    def code(self, r):
        for col in ('G', 'H'):
            m = CODE_RE.search(str(self.S.formula(r, col) or ''))
            if m:
                return m.group(0)
        return None

    def ref(self, r, col, prefix):
        f = self.S.formula(r, col)
        if f is None:
            return None
        m = re.fullmatch(r'=' + re.escape(prefix) + r'(\d+)', str(f).replace('$', ''))
        return int(m.group(1)) if m else None

    def line_formula(self, r):
        return re.fullmatch(rf'=I{r}\+U{r}', str(self.S.formula(r, 'V') or '').replace(' ', '')) is not None

    def scan(self):
        S = self.S
        for col in ('J', 'N', 'O', 'P', 'Q', 'R'):
            for r in range(FIRST_ROW, LAST_ROW + 1):
                if dec(S.value(r, col)) != 0:
                    self.err(f'SAAO {col}{r} is non-zero; the seed only reads SB-1..SB-3 (K..M) and no NTA')
        for r in range(11, self.R.f.max_row + 1):
            for c in range(CI('AG'), CI('BI')):
                v = self.R.v.cell(r, c).value
                if isinstance(v, (int, float)) and v != 0:
                    self.err(f'RealignmentReversion {CL(c)}{r} is non-zero; only G..AF are read')
        for r in range(11, self.O.f.max_row + 1):
            for col in ('S',):
                v = self.O.value(r, col)
                if isinstance(v, (int, float)) and v != 0:
                    self.err(f'OBLIGATION {col}{r} is non-zero; only G..R are read')

        self.rows = {}
        for r in range(FIRST_ROW, LAST_ROW + 1):
            text, indent = self.text(r)
            indent = INDENT_OVERRIDES.get(r, indent)
            vf = S.formula(r, 'V')
            is_line_f = self.line_formula(r)
            total = False
            if not is_line_f and is_formula(vf):
                total = True
            elif not is_line_f and text and TOTAL_TEXT_RE.match(text):
                total = True
            self.rows[r] = dict(text=text, indent=indent, total=total, line_f=is_line_f,
                                x=self.ref(r, 'X', 'OBLIGATION!T'),
                                t=self.ref(r, 'T', 'RealignmentReversion!BI'))
            for col, pre in (('X', 'OBLIGATION'), ('T', 'RealignmentReversion')):
                f = S.formula(r, col)
                if not total and is_formula(f) and pre in f and self.rows[r][col.lower()] is None:
                    self.err(f'SAAO {col}{r} has an unexpected formula {f!r}')
        for r, info in self.rows.items():
            if not info['total']:
                info['fig'] = self.read_figures(r)

        # Row kinds. Figures decide first: a row with any figure is a line.
        last_line_indent = None
        rows = list(range(FIRST_ROW, LAST_ROW + 1))
        for i, r in enumerate(rows):
            info = self.rows[r]
            if info['total']:
                info['kind'] = 'total'
                last_line_indent = None
                continue
            fig = info['fig']
            if fig['any']:
                if not (info['line_f'] or info['x']):
                    self.err(f'SAAO row {r} carries figures but has no line formulas')
                if not info['text']:
                    self.err(f'SAAO row {r} carries figures but has no description')
                kind = 'line'
            elif not info['text']:
                kind = None
            elif (last_line_indent is not None and info['indent'] > last_line_indent
                  and not self.deeper_next(rows, i)):
                kind = 'note'                       # descriptive row under a line (e.g. Greening Program)
            elif self.deeper_next(rows, i):
                kind = 'heading'                    # indented rows follow
            elif info['line_f']:
                kind = 'line'                       # a zero line (e.g. CNA / SRI with no amount yet)
            else:
                kind = 'heading'
            kind = KIND_OVERRIDES.get(r, kind)
            if kind == 'group':
                kind = 'heading'
            if kind == 'line':
                last_line_indent = info['indent']
            elif kind == 'heading':
                last_line_indent = None
            info['kind'] = kind

    def deeper_next(self, rows, i):
        nxt = self.next_content(rows, i)
        return nxt is not None and not self.rows[nxt]['total'] and \
            (self.rows[nxt]['indent'] or 0) > self.rows[rows[i]]['indent']

    def next_content(self, rows, i):
        for r in rows[i + 1:]:
            info = self.rows[r]
            if info['total'] or info['text']:
                return r
        return None

    def read_figures(self, r):
        S, O, R = self.S, self.O, self.R
        info = self.rows[r]
        fig = dict(appropriation=dec(S.value(r, 'I')), sb={}, rr={}, months={}, extras={})
        for n, col in SB_COLS.items():
            a = dec(S.value(r, col))
            if a != 0:
                fig['sb'][n] = a
        t = info['t']
        if t is not None:
            for col in RR_COLS:
                a = dec(R.value(t, col))
                if a != 0:
                    fig['rr'][col] = a
        x = info['x']
        if x is not None:
            for m, col in enumerate(OBL_MONTH_COLS, start=1):
                raw = O.formula(x, col)
                cached = O.value(x, col)
                if raw is None:
                    continue
                if is_formula(raw):
                    expr = raw[1:].replace(' ', '').replace(',', '')
                    if re.fullmatch(r'-?\d+(\.\d+)?([+-]\d+(\.\d+)?)*', expr):
                        terms = re.findall(r'[+-]?\d+(?:\.\d+)?', expr)
                        if any('.' in t_ and len(t_.split('.')[1]) > 2 for t_ in terms):
                            self.warnings.append(f'OBLIGATION {col}{x}: term with >2 decimals in {raw}')
                        amount = sum(Decimal(t_) for t_ in terms).quantize(CENT, ROUND_HALF_UP)
                        if amount != dec(cached):
                            self.err(f'OBLIGATION {col}{x}: {raw} evaluates to {amount}, cached {cached}')
                        entry = expr if len(terms) > 1 else None
                    else:
                        self.warnings.append(f'OBLIGATION {col}{x}: formula {raw!r} stored as its value')
                        amount, entry = dec(cached), None
                elif isinstance(raw, str) and numz(raw) == 0:
                    if raw.strip():
                        self.warnings.append(f'OBLIGATION {col}{x}: text {raw!r} ignored')
                    continue
                else:
                    amount, entry = dec(raw), None
                if amount != 0 or entry is not None:
                    fig['months'][m] = (amount, entry)
        for field, col in EXTRA_COLS.items():
            raw = S.formula(r, col)
            cached = S.value(r, col)
            if raw is None or (isinstance(raw, str) and raw.strip() == ''):
                continue
            if isinstance(cached, str) and cached.startswith('#'):
                self.warnings.append(f'SAAO {col}{r}: formula error {cached} skipped')
                continue
            if is_formula(raw):
                if re.fullmatch(rf'=Z{r}', raw.replace(' ', '')):
                    fig['extras'][field] = ('balance', None)
                else:
                    # e.g. AD497 '=AB497' copies the obligation % into a CNA column: a formula error.
                    self.warnings.append(f'SAAO {col}{r}: formula {raw!r} (cached {cached}) is not a tag; skipped')
                    self.skipped_extras.add((r, field))
            elif isinstance(raw, (int, float)):
                if raw != 0:
                    fig['extras'][field] = ('amount', dec(raw))
            else:
                self.warnings.append(f'SAAO {col}{r}: text {raw!r} ignored')
        if fig['extras'] and not (info['text'] or info['line_f']):
            self.warnings.append(f'SAAO row {r}: CNA/Savings tag on an empty row ignored')
            fig['extras'] = {}
        fig['any'] = bool(fig['appropriation'] or fig['sb'] or fig['rr'] or fig['months'] or fig['extras'])
        return fig

    # ---------------------------------------------------------------- blocks
    def parse_blocks(self):
        """Every total row's V formula -> block of rows it adds up."""
        self.blocks = {}
        for r in range(FIRST_ROW, LAST_ROW + 1):
            info = self.rows[r]
            if info['kind'] != 'total':
                continue
            f = str(self.S.formula(r, 'V') or '').replace(' ', '').replace('$', '')
            if not f.startswith('='):
                self.err(f'total row {r} ({info["text"]}) has no V formula')
                continue
            if re.sub(r'V\d+(:V\d+)?|SUM\(|\)|[+,=]', '', f):
                self.err(f'total row {r}: unsupported formula {f}')
                continue
            ranges, refs = [], []
            for a, b in re.findall(r'V(\d+)(?::V(\d+))?', f):
                if b:
                    ranges.append((int(a), int(b)))
                else:
                    refs.append(int(a))
            if ranges and refs:
                self.err(f'total row {r}: mixes ranges and references: {f}')
            if len(ranges) > 1:
                self.err(f'total row {r}: several ranges: {f}')
            self.blocks[r] = dict(row=r, ranges=ranges, refs=refs)
        for r, b in self.blocks.items():
            for a, e in b['ranges']:
                for q in range(a, e + 1):
                    if self.rows.get(q, {}).get('kind') == 'total':
                        self.err(f'total row {r}: its range contains total row {q} (double count)')
        self.first_cache = {}

    def first(self, r):
        """First content row of a block (or the row itself for a referenced line)."""
        if r not in self.blocks:
            return r
        b = self.blocks[r]
        if b['ranges']:
            return min(a for a, _ in b['ranges'])
        return min(self.first(q) for q in b['refs'])

    # ---------------------------------------------------------------- tree
    def add(self, key, kind, label, parent, indent=0, code=None, cls=None, total_label=None,
            synthetic=False):
        if key in self.nodes:
            self.err(f'row {key} placed twice')
        n = Node(key, kind, label, parent, indent, code, cls, total_label, synthetic)
        self.nodes[key] = n
        self.order.append(n)
        return n

    def heads_in(self, a, b):
        return [r for r in range(a, b) if self.rows[r]['kind'] == 'heading']

    def heading_node(self, r, parent, total_label=None):
        info = self.rows[r]
        cls = classification_of(info['text'])
        return self.add(r, 'classification' if cls else 'group', info['text'], parent,
                        info['indent'], cls=cls, total_label=total_label)

    def build(self):
        grand_i = str(self.S.formula(GRAND_ROW, 'I')).replace(' ', '')
        grand_v = str(self.S.formula(GRAND_ROW, 'V')).replace(' ', '')
        comps = [int(x) for x in re.findall(r'I(\d+)', grand_i)]
        if comps != [int(x) for x in re.findall(r'V(\d+)', grand_v)]:
            self.err('grand total I and V formulas list different rows')
        self.components = comps
        self.top = []
        prev = FIRST_ROW - 1
        for c in comps:
            if c not in self.blocks:              # stand-alone top-level line
                if self.rows[c]['kind'] != 'line':
                    self.err(f'grand-total component {c} is neither a total nor a line')
                for q in range(prev + 1, c):
                    self.check_skipped(q, f'before top-level line {c}')
                self.place_line(c, None)
                self.top.append(self.nodes[c])
                prev = c
                continue
            heads = self.heads_in(prev + 1, self.first(c))
            total_text = self.rows[c]['text']
            if c in SECTION_LABEL_FROM_TOTAL:
                label = re.sub(r'^TOTAL\s+APPROPRIATIONS?\s*-?\s*', '', total_text, flags=re.I)
                section = self.add(c, 'section', label, None, total_label=total_text, synthetic=True)
                wrappers = []
                for h in heads:
                    if classification_of(self.rows[h]['text']):
                        wrappers.append(h)
                    else:
                        break
                cursor = (wrappers[-1] + 1) if wrappers else prev + 1
            else:
                owner = next((h for h in heads if not classification_of(self.rows[h]['text'])), None)
                if owner is None:
                    self.err(f'section ending at row {c} has no heading')
                    prev = c
                    continue
                wrappers = [h for h in heads if h < owner]
                for q in range(prev + 1, owner):
                    if q not in wrappers:
                        self.check_skipped(q, f'before section heading {owner}')
                section = self.add(owner, 'section', self.rows[owner]['text'], None,
                                   self.rows[owner]['indent'], total_label=total_text)
                cursor = owner + 1
            self.owners[c] = section
            container = section
            for w in wrappers:
                container = self.heading_node(w, container)
            self.top.append(section)
            self.fill(c, cursor, container)
            prev = c

    def ancestor_class(self, node):
        while node is not None:
            if node.kind == 'classification':
                return node.cls
            node = node.parent
        return None

    def fill(self, b_row, cursor, container):
        """Place the content of block b_row (rows from cursor on) under container."""
        b = self.blocks[b_row]
        if b['ranges']:
            a, e = b['ranges'][0]
            self.process(cursor, e, container, covered_from=a)
            for q in range(e + 1, b_row):
                self.check_skipped(q, f'after the range of total row {b_row}')
            return
        for ref in b['refs']:
            if ref < cursor:
                self.err(f'total row {b_row}: reference V{ref} is out of order')
                continue
            if ref not in self.blocks and self.rows[ref]['kind'] is None:
                self.warnings.append(f'total row {b_row}: V{ref} is an empty row (adds 0)')
                continue
            if ref not in self.blocks:            # a line referenced directly
                if self.rows[ref]['kind'] != 'line':
                    self.err(f'total row {b_row}: V{ref} is neither a total nor a line')
                self.process(cursor, ref, container, covered_from=ref)
                cursor = ref + 1
            else:
                container = self.assign(ref, cursor, container)
                cursor = ref + 1
        for q in range(cursor, b_row):
            self.check_skipped(q, f'between the parts of total row {b_row}')

    def assign(self, b_row, cursor, container):
        """Choose the heading that owns block b_row; returns the container for later siblings."""
        heads = self.heads_in(cursor, self.first(b_row))
        total_text = self.rows[b_row]['text']
        want = class_of_total(total_text) if self.ancestor_class(container) is None else None
        owner = None
        if want:
            cls_heads = [h for h in heads if classification_of(self.rows[h]['text'])]
            if cls_heads:
                owner = cls_heads[0]
                for h in heads:                    # headings above the classification wrap the rest
                    if h >= owner:
                        break
                    container = self.heading_node(h, container)
            else:
                node = self.add(b_row, 'classification', CLASS_NAMES[want], container,
                                cls=want, total_label=total_text, synthetic=True)
                self.owners[b_row] = node
                self.fill(b_row, cursor, node)
                return container
        elif heads:
            owner = heads[0]
        elif self.blocks[b_row]['ranges'] and self.rows[self.first(b_row)]['kind'] == 'heading':
            owner = self.first(b_row)              # the range starts with its own heading
        if owner is None:
            self.fill(b_row, cursor, container)
            return container
        node = self.heading_node(owner, container, total_label=total_text)
        self.owners[b_row] = node
        self.fill(b_row, owner + 1, node)
        return container

    def place_line(self, r, parent):
        info = self.rows[r]
        n = self.add(r, 'line', info['text'] or '', parent, info['indent'] or 0, code=self.code(r))
        n.fig = info['fig']
        return n

    def check_skipped(self, r, where):
        info = self.rows.get(r)
        if not info or info['kind'] in (None, 'total'):
            return
        if info['kind'] == 'line' and info['fig']['any']:
            self.err(f'row {r} ({info["text"]}) has figures but is outside every total ({where})')
        elif info['kind'] in ('heading', 'note', 'line'):
            self.warnings.append(f'row {r} ({info["text"]}) {info["kind"]} without figures dropped ({where})')

    def chain(self, node, base):
        out = []
        while node is not None and node is not base:
            out.append(node)
            node = node.parent
        if node is None:
            self.err('override parent outside its block')
            return []
        return list(reversed(out))

    def process(self, start, end, base, covered_from):
        """Indentation-driven placement of rows start..end under base."""
        stack = []
        last_heading = False
        last_line = None
        for r in range(start, end + 1):
            info = self.rows[r]
            kind = info['kind']
            if kind is None:
                continue
            if kind == 'total':
                self.err(f'unexpected total row {r} inside a block')
                continue
            if kind == 'line' and r < covered_from and not info['fig']['any'] and info['text']:
                kind = 'heading'                    # a figure-free row above the summed range heads it
            if kind in ('line', 'note') and r < covered_from:
                self.check_skipped(r, f'before the range starting at {covered_from}')
                continue
            indent = info['indent'] or 0
            if r in PARENT_OVERRIDES:
                stack = [n for n in self.chain(self.nodes[PARENT_OVERRIDES[r]], base)]
                parent = stack[-1] if stack else base
                if kind == 'heading':
                    node = self.heading_node(r, parent)
                    stack.append(node)
                    last_heading = True
                elif kind == 'line':
                    last_line = self.place_line(r, parent)
                    last_heading = False
                else:
                    self.add(r, 'note', info['text'], parent, indent)
                continue
            if kind == 'heading':
                is_cls = classification_of(info['text']) is not None
                caps = info['text'] == info['text'].upper()
                if last_heading and stack and indent >= stack[-1].indent and \
                        not (is_cls and stack[-1].kind == 'classification'):
                    pass                            # directly under the heading just opened
                else:
                    while stack and stack[-1].indent > indent:
                        stack.pop()
                    if stack and stack[-1].indent == indent:
                        stack.pop()
                        while stack and stack[-1].indent == indent and stack[-1].caps and caps:
                            stack.pop()
                    while is_cls and stack and stack[-1].kind == 'classification':
                        stack.pop()
                if stack and stack[-1].child_indent is None:
                    stack[-1].child_indent = indent
                node = self.heading_node(r, stack[-1] if stack else base)
                stack.append(node)
                last_heading = True
            elif kind == 'line':
                if not last_heading:
                    while stack and (stack[-1].indent > indent or
                                     (stack[-1].child_indent is not None and indent < stack[-1].child_indent)):
                        stack.pop()
                if stack and stack[-1].child_indent is None:
                    stack[-1].child_indent = indent
                last_line = self.place_line(r, stack[-1] if stack else base)
                last_heading = False
            elif kind == 'note':
                parent = last_line.parent if last_line is not None else (stack[-1] if stack else base)
                self.add(r, 'note', info['text'], parent, indent)
                last_heading = False

    # ---------------------------------------------------------------- checks
    def line_calc(self, n):
        f = n.fig
        supp = sum(f['sb'].values(), Decimal(0))
        rr = sum(f['rr'].values(), Decimal(0))
        adjusted = f['appropriation'] + supp + rr
        obligation = sum((a for a, _ in f['months'].values()), Decimal(0))
        bal = adjusted - obligation
        ex = {k: (bal if mode == 'balance' else amt) for k, (mode, amt) in f['extras'].items()}
        return dict(appropriation=f['appropriation'], supplemental=supp, realignment=rr,
                    adjusted=adjusted, obligation=obligation, balance=bal,
                    sb={k: v for k, v in f['sb'].items()}, extras=ex,
                    savings=ex.get('savings', Decimal(0)), savings_half=half(ex.get('savings', Decimal(0))))

    def check(self):
        S = self.S
        lines = [n for n in self.order if n.kind == 'line']
        # every row with figures is a line node
        for r, info in self.rows.items():
            if info['kind'] == 'line' and info['fig']['any'] and r not in self.nodes:
                self.err(f'row {r} ({info["text"]}) has figures but was not placed')
        # per line against the workbook's cached values
        for n in lines:
            c = self.line_calc(n)
            r = n.key
            for col, mine in (('V', c['adjusted']), ('X', c['obligation'])):
                if dec(S.value(r, col)) != mine:
                    self.err(f'line {r} {col}: workbook {dec(S.value(r, col))}, seed {mine}')
            for field, col in EXTRA_COLS.items():
                wb = S.value(r, col)
                if isinstance(wb, str) and wb.startswith('#') or (r, field) in self.skipped_extras:
                    continue
                if dec(wb) != c['extras'].get(field, Decimal(0)):
                    self.err(f'line {r} {col}: workbook {dec(wb)}, seed {c["extras"].get(field, 0)}')
        # RR / OBLIGATION rows: every non-zero source row is read by exactly one line
        rr_users, ob_users = {}, {}
        for n in lines:
            info = self.rows[n.key]
            if info['t'] is not None:
                rr_users.setdefault(info['t'], []).append(n.key)
            if info['x'] is not None:
                ob_users.setdefault(info['x'], []).append(n.key)
        for r in range(FIRST_ROW, LAST_ROW + 1):
            rr_nz = any(numz(self.R.value(r, col)) != 0 for col in RR_COLS)
            if rr_nz and self.rows[r]['kind'] != 'total' and len(rr_users.get(r, [])) != 1:
                self.err(f'RealignmentReversion row {r} is non-zero and read by lines {rr_users.get(r, [])}')
            ob_nz = any(numz(self.O.value(r, col)) != 0 for col in OBL_MONTH_COLS)
            if ob_nz and self.rows[r]['kind'] != 'total' and len(ob_users.get(r, [])) != 1:
                self.err(f'OBLIGATION row {r} is non-zero and read by lines {ob_users.get(r, [])}')
        # every block owned by a node: subtree equals the total row
        for b_row in self.blocks:
            owner = self.owners.get(b_row)
            if owner is None:
                continue
            got = sum((self.line_calc(n)['adjusted'] for n in owner.lines()), Decimal(0))
            if got != dec(S.value(b_row, 'V')):
                self.err(f'block {b_row} ({self.rows[b_row]["text"]}) owned by "{owner.label}": '
                         f'workbook {dec(S.value(b_row, "V"))}, seed {got}')
        # sections
        self.section_report = []
        for c, node in zip(self.components, self.top):
            calc = [self.line_calc(n) for n in node.lines()]
            for col, key in (('I', 'appropriation'), ('V', 'adjusted'), ('X', 'obligation')):
                got = sum((x[key] for x in calc), Decimal(0))
                want = dec(S.value(c, col))
                if got != want:
                    self.err(f'section row {c} ({node.label}) {col}: workbook {want}, seed {got}')
            self.section_report.append((c, node.label, dec(S.value(c, 'V'))))
        # grand totals
        calc = [self.line_calc(n) for n in lines]
        sums = {
            'I': sum((x['appropriation'] for x in calc), Decimal(0)),
            'K': sum((x['sb'].get(1, 0) for x in calc), Decimal(0)),
            'L': sum((x['sb'].get(2, 0) for x in calc), Decimal(0)),
            'M': sum((x['sb'].get(3, 0) for x in calc), Decimal(0)),
            'S': sum((x['supplemental'] for x in calc), Decimal(0)),
            'T': sum((x['realignment'] for x in calc), Decimal(0)),
            'V': sum((x['adjusted'] for x in calc), Decimal(0)),
            'X': sum((x['obligation'] for x in calc), Decimal(0)),
            'Z': sum((x['balance'] for x in calc), Decimal(0)),
            'AE': sum((x['savings'] for x in calc), Decimal(0)),
            # The workbook's AF total adds unrounded AE/2 per line, i.e. half the summed savings.
            'AF': (sum((x['savings'] for x in calc), Decimal(0)) / 2).quantize(CENT, ROUND_HALF_UP),
        }
        self.half_lines = sum((x['savings_half'] for x in calc), Decimal(0))
        self.odd_savings = [(n.key, x['savings']) for n, x in zip(lines, calc)
                            if x['savings'] != x['savings_half'] * 2]
        self.grand = sums
        for col, got in sums.items():
            want = dec(S.value(GRAND_ROW, col))
            if got != want:
                self.err(f'grand total {col}: workbook {want}, seed {got}')

    # ---------------------------------------------------------------- movements
    def movements(self):
        out = []
        for n, col in SB_COLS.items():
            out.append(dict(col=f'SB{n}', kind='supplemental', sb=n, ref=None,
                            title=f'Supplemental Budget No. {n}', status='approved', note=None,
                            lines=[(ln.key, ln.fig['sb'][n]) for ln in self.order
                                   if ln.kind == 'line' and n in ln.fig['sb']]))
        for col in RR_COLS:
            lines = [(ln.key, ln.fig['rr'][col]) for ln in self.order
                     if ln.kind == 'line' and col in ln.fig['rr']]
            if not lines:
                continue
            title = collapse(self.R.formula(RR_HEADER_ROW, col) or '')
            if not title:
                self.err(f'RealignmentReversion column {col} has amounts but no header')
            u = title.upper()
            if 'SOURCE FOR SB' in u:
                kind = 'reversion'
            elif 'AUGMENTATION' in u:
                kind = 'augmentation'
            elif 'REALIGNMENT' in u or 'RESOLUTION' in u:
                kind = 'realignment'
            else:
                kind = 'other'
                self.warnings.append(f'column {col} "{title}" has no recognised kind; stored as other')
            m = re.search(r'\d{3}-\d{4}', title)
            ref = m.group(0) if m else None
            note = self.R.formula(RR_NOTE_ROW, col)
            note = collapse(note) if note is not None and collapse(note) else None
            if note is None:
                status = 'approved'
            elif 'CANCEL' in note.upper():
                status = 'cancelled'
            elif 'WAITING FOR APPROVAL' in note.upper() or 'FOR RESOLUTION' in note.upper():
                status = 'pending'
            else:
                self.err(f'column {col}: unrecognised status note {note!r}')
                status = 'pending'
            out.append(dict(col=col, kind=kind, sb=None, ref=ref, title=title, status=status,
                            note=note, lines=lines))
        for col in RR_COLS:
            note = self.R.formula(RR_NOTE_ROW, col)
            if note and collapse(note) and not any(m['col'] == col for m in out):
                self.warnings.append(f'column {col} ({collapse(note)}) has no amounts; no movement created')
        return out

    # ---------------------------------------------------------------- output
    def outline(self, out):
        def walk(n, d):
            extra = ''
            if n.kind == 'line':
                extra = f'  [{self.line_calc(n)["adjusted"]}]'
            tl = f'  <{n.total_label}>' if n.total_label else ''
            syn = ' *' if n.synthetic else ''
            out.write(f'{"  " * d}{n.key} {n.kind[:4]} {n.label[:70]}{syn}{tl}{extra}\n')
            for c in n.children:
                walk(c, d + 1)
        for n in self.top:
            walk(n, 0)


def q(s):
    if s is None:
        return 'null'
    return "'" + str(s).replace("'", "''") + "'"


def num(d):
    return f'{d:.2f}'


def emit(seed, out):
    S = seed.S
    header = [collapse(S.formula(r, 'A')) for r in HEADER_ROWS]
    cert_name = collapse(S.formula(CERTIFIER_ROWS[0], 'X'))
    cert_title = collapse(S.formula(CERTIFIER_ROWS[1], 'X'))
    nodes = seed.order
    lines = [n for n in nodes if n.kind == 'line']
    sort = {}
    for n in nodes:
        siblings = n.parent.children if n.parent else seed.top
        sort[n.key] = 10 * (siblings.index(n) + 1)
    mvs = seed.movements()
    mlines = [(m['col'], r, a) for m in mvs for r, a in m['lines']]
    obl = [(n.key, m, a, e) for n in lines for m, (a, e) in sorted(n.fig['months'].items())]
    extras = [(n.key, f, mode, a) for n in lines for f, (mode, a) in n.fig['extras'].items()]

    w = out.write
    w('-- GENERATED once by scripts/saao/gen_seed.py from the Budget Office workbook\n')
    w(f'-- "{SOURCE_NAME}" (October 2026). Do not edit by hand. Spec §12, "One-time seed".\n')
    w('-- Loads the FY2026 General Fund SAAO book: structure, account codes, annual appropriations,\n')
    w('-- SB-1..SB-3, one movement per RealignmentReversion column, monthly obligations and\n')
    w('-- CNA/Savings tags. Idempotent: does nothing when an FY2026 GF book already exists.\n')
    w('-- Verified by the generator (every line, section total and the grand total row 1364) and by\n')
    w('-- supabase/tests/19_saao_seed.test.sql. Keys in the temp tables are workbook row numbers\n')
    w('-- (SAAO sheet); synthesized headings use their total row number.\n')
    w('do $seed$\n')
    w('declare\n  v_fund uuid;\n  v_book uuid;\n  v_count int;\nbegin\n')
    w('  if exists (select 1 from pr_tracker.saao_books b join pr_tracker.saao_funds f on f.id = b.fund_id\n')
    w("             where b.fiscal_year = 2026 and f.code = 'GF') then\n    return;\n  end if;\n\n")
    w("  select id into v_fund from pr_tracker.saao_funds where code = 'GF';\n")
    w('  insert into pr_tracker.saao_books (fiscal_year, fund_id, title, header_lines, certifier_name, certifier_title)\n')
    hl = '[' + ', '.join('"' + h.replace('"', '\\"') + '"' for h in header) + ']'
    w(f"  values (2026, v_fund, 'FY 2026 General Fund', {q(hl)}::jsonb, {q(cert_name)}, {q(cert_title)})\n")
    w('  returning id into v_book;\n\n')

    w('  create temp table _seed_node (wb_row int primary key, id uuid not null default gen_random_uuid()) on commit drop;\n')
    keys = [n.key for n in nodes]
    w('  insert into _seed_node (wb_row) select unnest(array[\n')
    for i in range(0, len(keys), 20):
        w('    ' + ', '.join(str(k) for k in keys[i:i + 20]) + (',\n' if i + 20 < len(keys) else '\n'))
    w('  ]);\n\n')

    max_depth = max(n.depth() for n in nodes)
    for d in range(max_depth + 1):
        level = [n for n in nodes if n.depth() == d]
        if not level:
            continue
        w(f'  -- tree depth {d}\n')
        w('  insert into pr_tracker.saao_nodes (id, book_id, parent_id, sort_order, kind, label, account_code,\n')
        w('                                     classification_id, total_label, appropriation)\n')
        w('  select n.id, v_book, p.id, v.sort_order, v.kind, v.label, v.account_code, c.id, v.total_label, v.appropriation\n')
        w('  from (values\n')
        rows = []
        for i, n in enumerate(level):
            appr = n.fig['appropriation'] if n.kind == 'line' else Decimal(0)
            parent = n.parent.key if n.parent else None
            t = (lambda typ: f'::{typ}') if i == 0 else (lambda typ: '')   # first row fixes the types
            rows.append(f'    ({n.key}{t("int")}, {parent if parent is not None else "null"}{t("int")}, '
                        f'{sort[n.key]}{t("int")}, {q(n.kind)}{t("text")}, {q(n.label)}{t("text")}, '
                        f'{q(n.code)}{t("text")}, {q(n.cls if n.kind == "classification" else None)}{t("text")}, '
                        f'{q(n.total_label)}{t("text")}, {num(appr)}{t("numeric")})')
        w(',\n'.join(rows) + '\n')
        w('  ) as v (wb_row, parent_row, sort_order, kind, label, account_code, class_code, total_label, appropriation)\n')
        w('  join _seed_node n on n.wb_row = v.wb_row\n')
        w('  left join _seed_node p on p.wb_row = v.parent_row\n')
        w('  left join pr_tracker.saao_classifications c on c.code = v.class_code;\n\n')
    w('  select count(*) into v_count from pr_tracker.saao_nodes where book_id = v_book;\n')
    w(f"  if v_count <> {len(nodes)} then raise exception 'seed: expected {len(nodes)} nodes, got %', v_count; end if;\n\n")

    w('  create temp table _seed_mv (col text primary key, id uuid not null default gen_random_uuid()) on commit drop;\n')
    w('  insert into _seed_mv (col) values ' + ', '.join(f"({q(m['col'])})" for m in mvs) + ';\n')
    w('  insert into pr_tracker.saao_movements (id, book_id, kind, sb_number, reference_no, title, effective_date,\n')
    w('                                         status, status_note, remarks)\n')
    w(f"  select m.id, v_book, v.kind, v.sb_number, v.reference_no, v.title, date '{EFFECTIVE_DATE}', v.status,\n")
    w(f'         v.status_note, {q(REMARK)}\n')
    w('  from (values\n')
    rows = []
    for i, m in enumerate(mvs):
        sb = str(m['sb']) if m['sb'] else 'null'
        if i == 0:
            sb += '::int'
        rows.append(f"    ({q(m['col'])}, {q(m['kind'])}, {sb}, {q(m['ref'])}{'::text' if i == 0 else ''}, "
                    f"{q(m['title'])}, {q(m['status'])}, {q(m['note'])}{'::text' if i == 0 else ''})")
    w(',\n'.join(rows) + '\n')
    w('  ) as v (col, kind, sb_number, reference_no, title, status, status_note)\n')
    w('  join _seed_mv m on m.col = v.col;\n\n')

    w('  insert into pr_tracker.saao_movement_lines (movement_id, node_id, amount)\n')
    w('  select m.id, n.id, v.amount from (values\n')
    rows = [f'    ({q(c)}, {r}, {num(a)})' for c, r, a in mlines]
    w(',\n'.join(rows) + '\n')
    w('  ) as v (col, wb_row, amount)\n')
    w('  join _seed_mv m on m.col = v.col join _seed_node n on n.wb_row = v.wb_row;\n')
    w('  select count(*) into v_count from pr_tracker.saao_movement_lines ml\n')
    w('    join pr_tracker.saao_movements m on m.id = ml.movement_id where m.book_id = v_book;\n')
    w(f"  if v_count <> {len(mlines)} then raise exception 'seed: expected {len(mlines)} movement lines, got %', v_count; end if;\n\n")

    w('  insert into pr_tracker.saao_obligations (node_id, month, amount, entry)\n')
    w('  select n.id, v.month, v.amount, v.entry from (values\n')
    rows = []
    for i, (r, m, a, e) in enumerate(obl):
        rows.append(f'    ({r}, {m}{"::smallint" if i == 0 else ""}, {num(a)}, {q(e)}{"::text" if i == 0 else ""})')
    w(',\n'.join(rows) + '\n')
    w('  ) as v (wb_row, month, amount, entry)\n')
    w('  join _seed_node n on n.wb_row = v.wb_row;\n')
    w('  select count(*) into v_count from pr_tracker.saao_obligations o\n')
    w('    join pr_tracker.saao_nodes n on n.id = o.node_id where n.book_id = v_book;\n')
    w(f"  if v_count <> {len(obl)} then raise exception 'seed: expected {len(obl)} obligations, got %', v_count; end if;\n\n")

    w('  insert into pr_tracker.saao_line_extras (node_id, field, mode, amount)\n')
    w('  select n.id, v.field, v.mode, v.amount from (values\n')
    rows = []
    for i, (r, f, mode, a) in enumerate(extras):
        rows.append(f'    ({r}, {q(f)}, {q(mode)}, {num(a) if a is not None else "null"}{"::numeric" if i == 0 else ""})')
    w(',\n'.join(rows) + '\n')
    w('  ) as v (wb_row, field, mode, amount)\n')
    w('  join _seed_node n on n.wb_row = v.wb_row;\n')
    w('  select count(*) into v_count from pr_tracker.saao_line_extras e\n')
    w('    join pr_tracker.saao_nodes n on n.id = e.node_id where n.book_id = v_book;\n')
    w(f"  if v_count <> {len(extras)} then raise exception 'seed: expected {len(extras)} extras, got %', v_count; end if;\n\n")

    g = seed.grand
    after = (f"jsonb_build_object('source', {q(SOURCE_NAME)}, 'generator', 'scripts/saao/gen_seed.py', "
             f"'nodes', {len(nodes)}, 'lines', {len(lines)}, 'movements', {len(mvs)}, "
             f"'movement_lines', {len(mlines)}, 'obligations', {len(obl)}, 'extras', {len(extras)}, "
             f"'appropriation', {num(g['I'])}, 'adjusted', {num(g['V'])}, 'obligation', {num(g['X'])})")
    w('  insert into pr_tracker.saao_audit (book_id, actor_id, entity, entity_id, action, before, after)\n')
    w(f"  values (v_book, null, 'book', v_book, 'seed', null,\n          {after});\n")
    w('end\n$seed$;\n')


def main(argv):
    if len(argv) < 2:
        die('usage: gen_seed.py <workbook.xlsx> [--outline]')
    seed = Seed(argv[1])
    seed.scan()
    seed.parse_blocks()
    seed.build()
    seed.check()
    for wmsg in seed.warnings:
        print('warning: ' + wmsg, file=sys.stderr)
    for c, label, v in seed.section_report:
        print(f'section {c:>5} {v:>18,.2f}  {label}', file=sys.stderr)
    if '--outline' in argv:
        seed.outline(sys.stdout)
    if seed.errors:
        for e in seed.errors:
            print('MISMATCH: ' + e, file=sys.stderr)
        print(f'{len(seed.errors)} mismatch(es); no SQL written', file=sys.stderr)
        return 1
    g = seed.grand
    print('grand totals match row 1364: ' + ', '.join(f'{k}={v}' for k, v in g.items()), file=sys.stderr)
    print(f'note: sum of per-line 50% savings (R-B rounding) = {seed.half_lines}; lines with odd-centavo '
          f'savings: {seed.odd_savings}', file=sys.stderr)
    if '--outline' not in argv:
        emit(seed, sys.stdout)
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
